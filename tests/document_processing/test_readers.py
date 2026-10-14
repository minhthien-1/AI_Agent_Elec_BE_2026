import json
from pathlib import Path

import pymupdf
from docx import Document

from document_processing.cli import process_inputs
from document_processing.reader import read_document


def _create_pdf(path: Path, pages: int = 1) -> None:
    document = pymupdf.open()
    for page_number in range(1, pages + 1):
        page = document.new_page(width=300, height=400)
        page.insert_text((30, 20), "Elec-Agent manual")
        for line_number in range(4):
            page.insert_text(
                (30, 120 + line_number * 24),
                f"Technical content page {page_number} section {line_number} safety procedures.",
            )
        page.insert_text((30, 390), "Page footer")
    document.save(path)
    document.close()


def test_read_pdf_retains_physical_page_count_and_geometry(tmp_path: Path) -> None:
    source = tmp_path / "manual.pdf"
    _create_pdf(source, pages=2)

    document = read_document(source)

    assert document.file_type == "pdf"
    assert document.page_count == 2
    assert document.pages[0].width == 300
    assert document.pages[0].height == 400
    assert document.pages[0].blocks[0].bbox is not None
    assert "Technical content page 2" in document.raw_text


def test_read_docx_retains_heading_and_table(tmp_path: Path) -> None:
    source = tmp_path / "guide.docx"
    document = Document()
    document.add_heading("Safety heading", level=1)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Voltage"
    table.cell(0, 1).text = "220 V"
    document.save(source)

    extracted = read_document(source)

    assert extracted.file_type == "docx"
    assert [block.kind for block in extracted.pages[0].blocks] == [
        "heading",
        "table_row",
    ]
    assert "Safety heading" in extracted.raw_text
    assert "Voltage | 220 V" in extracted.raw_text


def test_read_html_retains_headings_and_ignores_scripts(tmp_path: Path) -> None:
    source = tmp_path / "guide.html"
    source.write_text(
        "<html><head><title>Guide</title></head><body>"
        "<h1>Safety heading</h1><p>Keep this content.</p>"
        "<script>ignore_this_script()</script></body></html>",
        encoding="utf-8",
    )

    document = read_document(source)

    assert document.metadata["title"] == "Guide"
    assert document.pages[0].blocks[0].kind == "heading"
    assert "Safety heading" in document.raw_text
    assert "Keep this content." in document.raw_text
    assert "ignore_this_script" not in document.raw_text


def test_read_txt_records_encoding(tmp_path: Path) -> None:
    source = tmp_path / "guide.txt"
    source.write_bytes("Tiếng Việt: an toàn điện".encode("utf-16"))

    document = read_document(source)

    assert document.file_type == "txt"
    assert document.metadata["encoding"] == "utf-16"
    assert document.raw_text == "Tiếng Việt: an toàn điện"


def test_cli_writes_txt_and_json_for_document(tmp_path: Path) -> None:
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    _create_pdf(source_dir / "manual.pdf")
    output_dir = tmp_path / "output"

    summary, failures = process_inputs(str(source_dir), str(output_dir))

    assert summary == {
        "discovered": 1,
        "processed": 1,
        "failed": 0,
        "warning_count": 0,
    }
    assert failures == []
    assert (output_dir / "manual.txt").is_file()
    payload = json.loads((output_dir / "manual.json").read_text(encoding="utf-8"))
    assert payload["page_count"] == 1
    assert payload["cleaning_stats"]["cleaned_char_count"] > 0
